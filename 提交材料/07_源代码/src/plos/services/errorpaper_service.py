"""错题本试卷导出服务。

把错题本中的错题按中小学标准试卷版式导出为 Word(.docx) / PDF：
- 顶部标题 + 生成日期 + 知识点范围
- 页眉「PLOS AI 个人学习操作系统」、页脚「第 X 页 / 共 Y 页」
- 题型自动分块（选择题 / 填空题 / 判断题 / 简答题），全程统一有序题号
- 选择题按 A、B、C、D 规范排版并预留答题空括号；简答题预留大片作答空白
- 勾选「带答案」时，答案与解析统一放在文末【参考答案与解析】，不穿插在题间

依赖说明（打包友好，均已在 requirements / 打包产物中，无新增重量级依赖）：
- Word：python-docx
- PDF：PyMuPDF（现有依赖，无需安装 Word 或 wkhtmltopdf 等外部工具）
"""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Sequence, Tuple

from ..db import Database, get_db
from ..utils.latex_clean import latex_to_unicode, split_fractions
from ..utils.logger import get_logger
from ..utils.math_latex import strip_math_delimiters
from ..utils.paths import get_attachments_dir, get_images_dir
from .errorbook_service import ErrorBookService
from .user_service import UserService

logger = get_logger("services.errorpaper_service")

#: 试卷主标题
PAPER_TITLE = "PLOS AI 错题练习卷"
#: 页眉文字
PAPER_HEADER = "PLOS AI 个人学习操作系统"

#: 默认排版选项
DEFAULT_OPTIONS: Dict[str, Any] = {
    "title": PAPER_TITLE,
    "with_answers": True,        # 是否附参考答案与解析
    "with_wrong_answer": False,  # 是否附「原错误作答记录」
    "blank_lines": 6,            # 简答/主观题预留作答空白行数
    "knowledge_label": "全部错题",
}

#: 题型分组（顺序固定，决定试卷分块顺序）
TYPE_BUCKETS: List[Tuple[str, str]] = [
    ("choice", "选择题"),
    ("fill", "填空题"),
    ("judge", "判断题"),
    ("short", "简答题"),
]

#: 题型识别关键词（错题本 question_type 是自由文本，需归并到四个分块）
_TYPE_KEYWORDS: List[Tuple[str, Tuple[str, ...]]] = [
    ("choice", ("选择", "单选", "多选", "choice")),
    ("judge", ("判断", "judge")),
    ("fill", ("填空", "fill")),
    (
        "short",
        (
            "简答", "解答", "答题", "主观", "论述", "开放", "应用",
            "计算", "证明", "作图", "short", "open",
        ),
    ),
]

#: 中文节序号
_CN_NUMBERS = ["一", "二", "三", "四", "五", "六", "七", "八"]

#: 题干内联选项：A. / A、/ A．/ A) / A）
_OPTION_SPLIT_RE = re.compile(r"(?:^|(?<=\s))([A-D])\s*[.、．)）]\s*")
#: Windows 文件名非法字符
_BAD_FILENAME_CHARS = re.compile(r'[\\/:*?"<>|\r\n\t]')


def sanitize_filename(name: str, fallback: str = "PLOS-错题试卷") -> str:
    """清掉 Windows 文件名非法字符，避免导出时因为文件名报错。"""
    cleaned = _BAD_FILENAME_CHARS.sub("", str(name or "")).strip().strip(".")
    return cleaned or fallback


def default_filename(prefix: str = "PLOS-错题试卷") -> str:
    """默认文件名：PLOS-错题试卷_YYYYMMDD。"""
    return f"{prefix}_{datetime.now().strftime('%Y%m%d')}"


def _extract_options(text: str) -> Tuple[str, List[str]]:
    """从题干文本里拆出内联的 A/B/C/D 选项。

    错题本没有独立的选项字段，选择题的选项通常直接写在题干里。
    只有从 A 开始且按序连续的选项才被认可，避免误伤普通文本。
    """
    if not text:
        return text or "", []
    matches = list(_OPTION_SPLIT_RE.finditer(text))
    if len(matches) < 2:
        return text, []

    sequential: List[Any] = []
    expected = ord("A")
    for m in matches:
        if ord(m.group(1)) == expected:
            sequential.append(m)
            expected += 1
        else:
            break
    if len(sequential) < 2:
        return text, []

    stem = text[: sequential[0].start()].rstrip()
    options: List[str] = []
    for index, m in enumerate(sequential):
        end = sequential[index + 1].start() if index + 1 < len(sequential) else len(text)
        options.append(f"{m.group(1)}. {text[m.end():end].strip()}")
    return stem, options


def _classify(error: Dict[str, Any]) -> str:
    """把错题的 question_type 归并到四个题型分块之一。"""
    raw = str(error.get("question_type") or "").strip().lower()
    if raw:
        for key, words in _TYPE_KEYWORDS:
            if any(w in raw for w in words):
                return key
    # 题型为空或无法识别：按题干是否内联选项兜底判断，否则按主观题处理
    if _extract_options(str(error.get("question") or ""))[1]:
        return "choice"
    return "short"


def _has_answer_bracket(text: str) -> bool:
    """题干是否已自带答题括号（（　）或 ( )）。"""
    return ("（" in text and "）" in text) or ("(" in text and ")" in text)


class ErrorPaperService:
    """错题本试卷导出业务服务（不含任何 PyQt 依赖，可在后台线程调用）。"""

    # 字号规范：标题二号 22pt、题目小四 12pt、解析五号 10.5pt
    _FONT_CN = "宋体"
    _FONT_EN = "Times New Roman"
    _SIZE_TITLE = 22.0
    _SIZE_QUESTION = 12.0
    _SIZE_ANALYSIS = 10.5
    _SIZE_HEADING = 14.0
    _LINE_SPACING = 1.5
    _MARGIN_V_CM = 2.54   # 上下页边距（标准）
    _MARGIN_H_CM = 3.18   # 左右页边距（标准）

    def __init__(
        self,
        db: Optional[Database] = None,
        errorbook_service: Optional[ErrorBookService] = None,
        user_service: Optional[UserService] = None,
    ):
        self.db = db or get_db()
        self.user_service = user_service
        self.errorbook_service = errorbook_service or ErrorBookService(
            db=self.db, user_service=user_service
        )

    def _user_id(self, user_id: Optional[int] = None) -> int:
        if user_id is not None:
            return user_id
        if self.user_service is not None:
            return self.user_service.get_current_user_id()
        return 0

    # ------------------------------------------------------------------
    # 选题
    # ------------------------------------------------------------------

    def list_knowledge_points(self, user_id: Optional[int] = None) -> List[str]:
        """读取错题库内全部知识点标签（供下拉筛选）。"""
        return self.errorbook_service.list_knowledge_points(user_id=user_id)

    def select_errors(
        self,
        scope: str = "all",
        knowledge_point: str = "",
        error_ids: Optional[Sequence[int]] = None,
        subject: str = "",
        user_id: Optional[int] = None,
    ) -> List[Dict[str, Any]]:
        """按选题范围取错题。

        scope:
            - "all"：全部错题
            - "knowledge_point"：按知识点筛选
            - "custom"：自定义勾选的题目（error_ids）
        """
        if scope == "custom":
            ids = [int(i) for i in (error_ids or [])]
            if not ids:
                return []
            uid = self._user_id(user_id)
            placeholders = ",".join("?" for _ in ids)
            rows = self.db.fetchall(
                f"SELECT * FROM error_book WHERE user_id = ? AND id IN ({placeholders})",
                (uid, *ids),
            )
            order = {eid: pos for pos, eid in enumerate(ids)}
            return sorted(rows, key=lambda r: order.get(int(r["id"]), len(order)))

        if scope == "knowledge_point":
            if not knowledge_point:
                return []
            return self.errorbook_service.list_errors(
                knowledge_point=knowledge_point,
                subject=subject or None,
                user_id=user_id,
            )

        return self.errorbook_service.list_errors(
            subject=subject or None, user_id=user_id
        )

    def list_errors_for_pick(
        self, subject: str = "", user_id: Optional[int] = None
    ) -> List[Dict[str, Any]]:
        """自定义选题用的候选列表（可按学科过滤）。"""
        return self.errorbook_service.list_errors(
            subject=subject or None, user_id=user_id
        )

    # ------------------------------------------------------------------
    # 试卷模型
    # ------------------------------------------------------------------

    def build_paper(
        self,
        errors: Sequence[Dict[str, Any]],
        options: Optional[Dict[str, Any]] = None,
        user_id: Optional[int] = None,
    ) -> Dict[str, Any]:
        """把错题整理成可直接排版的试卷模型（纯数据，便于单测）。"""
        opts = {**DEFAULT_OPTIONS, **(options or {})}
        wrong_answers = (
            self._load_wrong_answers([int(e["id"]) for e in errors], user_id)
            if opts.get("with_wrong_answer")
            else {}
        )

        buckets: Dict[str, List[Dict[str, Any]]] = {key: [] for key, _ in TYPE_BUCKETS}
        for error in errors:
            buckets[_classify(error)].append(dict(error))

        number = 0
        sections: List[Dict[str, Any]] = []
        for key, label in TYPE_BUCKETS:
            items = buckets.get(key) or []
            if not items:
                continue
            questions: List[Dict[str, Any]] = []
            for error in items:
                number += 1
                questions.append(
                    self._build_question(error, key, number, opts, wrong_answers)
                )
            sections.append({"key": key, "label": label, "questions": questions})

        knowledge_scope = str(opts.get("knowledge_label") or "全部错题")
        meta = (
            f"生成日期：{datetime.now().strftime('%Y-%m-%d')}　　"
            f"知识点范围：{knowledge_scope}　　"
            f"题量：{number} 题"
        )
        return {
            "title": str(opts.get("title") or PAPER_TITLE),
            "meta": meta,
            "knowledge_scope": knowledge_scope,
            "generated_at": datetime.now().strftime("%Y-%m-%d"),
            "total": number,
            "sections": sections,
            "with_answers": bool(opts.get("with_answers", True)),
            "with_wrong_answer": bool(opts.get("with_wrong_answer", False)),
            "answers": self._collect_answers(sections, bool(opts.get("with_answers", True))),
        }

    def _build_question(
        self,
        error: Dict[str, Any],
        bucket: str,
        number: int,
        opts: Dict[str, Any],
        wrong_answers: Dict[int, str],
    ) -> Dict[str, Any]:
        """整理单道题：题干、选项、作答空白、图片。

        这里保留原始文本（含 LaTeX），由各渲染器在写入时再转 Unicode，
        这样 \\frac 才能被排版成上下结构的真分数而不是内联的 a/b。
        """
        stem, inline_options = _extract_options(str(error.get("question") or ""))
        stem = stem.strip()
        if not latex_to_unicode(stem).strip():
            stem = "（题干缺失）"
        options: List[str] = []
        blank_lines = 0

        if bucket == "choice":
            options = [o.strip() for o in inline_options]
            if not _has_answer_bracket(stem):
                stem = f"{stem}（　　）"
        elif bucket == "judge":
            if not _has_answer_bracket(stem):
                stem = f"{stem}（　　）"
        elif bucket == "short":
            # 简答/主观题预留大片作答空白
            blank_lines = max(0, int(opts.get("blank_lines", 6) or 0))

        return {
            "number": number,
            "stem": stem,
            "options": options,
            "answer": str(error.get("answer") or "").strip(),
            "analysis": str(error.get("analysis") or "").strip(),
            "wrong_answer": str(wrong_answers.get(int(error["id"]), "")).strip(),
            "blank_lines": blank_lines,
            "image": self._resolve_image(error.get("image_path")),
            "knowledge_point": str(error.get("knowledge_point") or "").strip(),
        }

    @staticmethod
    def _collect_answers(
        sections: Sequence[Dict[str, Any]], with_answers: bool
    ) -> List[Dict[str, Any]]:
        """收集文末答案板块的数据（统一放最后，不穿插在题间）。"""
        if not with_answers:
            return []
        answers: List[Dict[str, Any]] = []
        for section in sections:
            for q in section["questions"]:
                answers.append(
                    {
                        "number": q["number"],
                        "answer": q["answer"] or "（未填写）",
                        "analysis": q["analysis"],
                        "wrong_answer": q.get("wrong_answer", ""),
                    }
                )
        return answers

    def _load_wrong_answers(
        self, error_ids: Sequence[int], user_id: Optional[int] = None
    ) -> Dict[int, str]:
        """取每条错题最近一次的主观题作答记录（原错误作答）。"""
        ids = [int(i) for i in error_ids if i is not None]
        if not ids:
            return {}
        uid = self._user_id(user_id)
        placeholders = ",".join("?" for _ in ids)
        try:
            rows = self.db.fetchall(
                "SELECT error_id, user_answer FROM subjective_grades "
                f"WHERE user_id = ? AND error_id IN ({placeholders}) "
                "ORDER BY created_at ASC, id ASC",
                (uid, *ids),
            )
        except Exception as e:  # 表缺失等异常不影响导出主流程
            logger.warning("Load wrong answers failed: %s", e)
            return {}
        result: Dict[int, str] = {}
        for row in rows:
            text = str(row.get("user_answer") or "").strip()
            if text:
                result[int(row["error_id"])] = text  # 升序遍历，后写覆盖 = 取最新
        return result

    @staticmethod
    def _resolve_image(raw: Any) -> Optional[Path]:
        """解析错题图片路径（绝对路径优先，其次 images / attachments 目录）。"""
        if not raw:
            return None
        text = str(raw).strip()
        if not text:
            return None
        candidate = Path(text)
        if candidate.is_absolute() and candidate.exists():
            return candidate
        for base in (get_images_dir(), get_attachments_dir()):
            path = base / text
            if path.exists():
                return path
        return None

    # ------------------------------------------------------------------
    # 导出入口
    # ------------------------------------------------------------------

    def export(
        self,
        errors: Sequence[Dict[str, Any]],
        output_path: Path,
        fmt: str = "docx",
        options: Optional[Dict[str, Any]] = None,
        user_id: Optional[int] = None,
        progress: Optional[Callable[[int, str], None]] = None,
    ) -> Path:
        """生成试卷文件，返回最终文件路径。"""
        if not errors:
            raise ValueError("没有可导出的题目，请先选择错题范围")

        def report(percent: int, message: str) -> None:
            if progress is not None:
                try:
                    progress(percent, message)
                except Exception:  # 进度回调异常不应中断导出
                    pass

        report(15, "正在整理题目…")
        paper = self.build_paper(errors, options, user_id=user_id)
        report(50, "正在生成文档…")

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        if fmt == "docx":
            self._render_docx(paper, output_path)
        elif fmt == "pdf":
            self._render_pdf(paper, output_path)
        else:
            raise ValueError(f"不支持的导出格式：{fmt}")
        report(100, "生成完成")
        logger.info("Exported error paper (%s) -> %s", fmt, output_path)
        return output_path

    # ------------------------------------------------------------------
    # Word 渲染（python-docx）
    # ------------------------------------------------------------------

    def _render_docx(self, paper: Dict[str, Any], path: Path) -> None:
        try:
            from docx import Document
            from docx.enum.text import WD_ALIGN_PARAGRAPH
        except ImportError as e:  # pragma: no cover - 依赖缺失时给友好提示
            raise RuntimeError("缺少 python-docx 依赖，无法导出 Word 试卷") from e

        document = Document()
        self._setup_docx_section(document)
        self._setup_docx_styles(document)
        self._write_docx_header_footer(document)

        # 标题（二号，居中）+ 生成日期 / 知识点范围（五号，居中）
        self._docx_paragraph(
            document, paper["title"], self._SIZE_TITLE,
            WD_ALIGN_PARAGRAPH.CENTER, bold=True,
        )
        self._docx_paragraph(
            document, paper["meta"], self._SIZE_ANALYSIS, WD_ALIGN_PARAGRAPH.CENTER
        )

        for index, section in enumerate(paper["sections"]):
            self._docx_paragraph(
                document,
                f"{_CN_NUMBERS[index]}、{section['label']}",
                self._SIZE_HEADING,
                WD_ALIGN_PARAGRAPH.LEFT,
                bold=True,
                space_before=12,
                space_after=6,
            )
            for q in section["questions"]:
                self._docx_question(document, q, WD_ALIGN_PARAGRAPH)

        if paper["answers"]:
            document.add_page_break()
            self._docx_paragraph(
                document, "参考答案与解析", self._SIZE_HEADING,
                WD_ALIGN_PARAGRAPH.CENTER, bold=True, space_after=8,
            )
            for item in paper["answers"]:
                self._docx_answer(document, item, WD_ALIGN_PARAGRAPH)

        document.save(str(path))

    def _setup_docx_section(self, document) -> None:
        from docx.shared import Cm

        section = document.sections[0]
        section.page_width = Cm(21.0)   # A4
        section.page_height = Cm(29.7)
        section.top_margin = Cm(self._MARGIN_V_CM)
        section.bottom_margin = Cm(self._MARGIN_V_CM)
        section.left_margin = Cm(self._MARGIN_H_CM)
        section.right_margin = Cm(self._MARGIN_H_CM)
        section.header_distance = Cm(1.5)
        section.footer_distance = Cm(1.75)

    def _setup_docx_styles(self, document) -> None:
        from docx.oxml.ns import qn
        from docx.shared import Pt

        style = document.styles["Normal"]
        style.font.name = self._FONT_EN
        style.font.size = Pt(self._SIZE_QUESTION)
        try:
            style.element.rPr.rFonts.set(qn("w:eastAsia"), self._FONT_CN)
        except Exception:
            pass
        try:
            style.paragraph_format.line_spacing = self._LINE_SPACING
            style.paragraph_format.space_after = Pt(0)
        except Exception:
            pass

    def _write_docx_header_footer(self, document) -> None:
        """页眉 = 系统名（带下边框）；页脚 = 第 X 页 / 共 Y 页（真实域）。"""
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        section = document.sections[0]

        header = section.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.text = ""
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._style_docx_run(hp.add_run(PAPER_HEADER), 9.0)
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "999999")
        borders.append(bottom)
        hp._p.get_or_add_pPr().append(borders)

        footer = section.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.text = ""
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._style_docx_run(fp.add_run("第 "), 9.0)
        self._add_docx_field(fp, "PAGE")
        self._style_docx_run(fp.add_run(" 页 / 共 "), 9.0)
        self._add_docx_field(fp, "NUMPAGES")
        self._style_docx_run(fp.add_run(" 页"), 9.0)

    def _style_docx_run(self, run, size: float, bold: bool = False) -> None:
        from docx.oxml.ns import qn
        from docx.shared import Pt

        run.font.name = self._FONT_EN
        run.font.size = Pt(size)
        run.font.bold = bold
        try:
            run._element.rPr.rFonts.set(qn("w:eastAsia"), self._FONT_CN)
        except Exception:
            pass

    @staticmethod
    def _add_docx_field(paragraph, instruction: str) -> None:
        """插入 Word 域（PAGE / NUMPAGES），实现自动页码。"""
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        run = paragraph.add_run()
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = f" {instruction} "
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.append(begin)
        run._r.append(instr)
        run._r.append(end)

    def _docx_paragraph(
        self,
        document,
        text: str,
        size: float,
        alignment,
        bold: bool = False,
        space_before: float = 0.0,
        space_after: float = 0.0,
        indent_cm: float = 0.0,
        auto_indent: bool = False,
    ):
        """添加一个试卷段落。

        文本里可能带 ``\\( \\)``/``\\[ \\]`` 公式标记：先剥掉标记，
        再把 ``\\frac`` 渲染成 Word 原生上下分数（OMML），其余部分转可读 Unicode，
        保证导出的文档里不会出现原始 LaTeX 代码。
        """
        from docx.shared import Cm, Pt

        paragraph = document.add_paragraph()
        paragraph.alignment = alignment
        if space_before:
            paragraph.paragraph_format.space_before = Pt(space_before)
        if space_after:
            paragraph.paragraph_format.space_after = Pt(space_after)
        if indent_cm:
            paragraph.paragraph_format.left_indent = Cm(indent_cm)
        elif auto_indent:
            # 长题干悬挂缩进，题号左对齐、续行缩进，符合试卷观感
            paragraph.paragraph_format.left_indent = Cm(0.75)
            paragraph.paragraph_format.first_line_indent = Cm(-0.75)

        self._add_runs_with_fractions(paragraph, text or "", size, bold)
        return paragraph

    def _add_runs_with_fractions(self, paragraph, text: str, size: float, bold: bool) -> None:
        """把文本写入段落：普通文字用 run，分数用 OMML 原生公式。"""
        plain = strip_math_delimiters(text)
        added = False
        for chunk, frac in split_fractions(plain):
            if chunk:
                cleaned = latex_to_unicode(chunk)
                if cleaned:
                    self._style_docx_run(paragraph.add_run(cleaned), size, bold)
                    added = True
            if frac:
                paragraph._p.append(self._omml_fraction(frac[0], frac[1]))
                added = True
        if not added:
            self._style_docx_run(paragraph.add_run(""), size, bold)

    @staticmethod
    def _omml_fraction(num: str, den: str):
        """构造 OMML 上下分数元素（Word 原生公式格式）。"""
        from xml.sax.saxutils import escape

        from docx.oxml import parse_xml

        num_text = latex_to_unicode(num).strip() or " "
        den_text = latex_to_unicode(den).strip() or " "
        omml = (
            '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'
            ' xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            "<m:f>"
            f'<m:num><m:r><m:t xml:space="preserve">{escape(num_text)}</m:t></m:r></m:num>'
            f'<m:den><m:r><m:t xml:space="preserve">{escape(den_text)}</m:t></m:r></m:den>'
            "</m:f></m:oMath>"
        )
        return parse_xml(omml)

    def _docx_question(self, document, q: Dict[str, Any], align) -> None:
        """写入一道题：题干 + 选项 + 作答空白 + 图片。"""
        self._docx_paragraph(
            document,
            f"{q['number']}．{q['stem']}",
            self._SIZE_QUESTION,
            align.LEFT,
            space_before=4,
            auto_indent=True,
        )
        for option in q.get("options") or []:
            self._docx_paragraph(
                document, option, self._SIZE_QUESTION, align.LEFT, indent_cm=0.75
            )
        image = q.get("image")
        if image is not None:
            try:
                from docx.shared import Cm

                picture = document.add_paragraph()
                picture.alignment = align.CENTER
                picture.add_run().add_picture(str(image), width=Cm(10))
            except Exception as e:  # 图片损坏 / 格式不支持时跳过，不中断导出
                logger.warning("Skip error image %s: %s", image, e)
        for _ in range(int(q.get("blank_lines") or 0)):
            blank = document.add_paragraph()
            blank.paragraph_format.line_spacing = self._LINE_SPACING

    def _docx_answer(self, document, item: Dict[str, Any], align) -> None:
        """写入文末答案板块的一条记录。"""
        self._docx_paragraph(
            document,
            f"{item['number']}．【答案】{item['answer']}",
            self._SIZE_ANALYSIS,
            align.LEFT,
            space_before=4,
        )
        if item.get("analysis"):
            self._docx_paragraph(
                document, f"【解析】{item['analysis']}",
                self._SIZE_ANALYSIS, align.LEFT, indent_cm=0.5,
            )
        if item.get("wrong_answer"):
            self._docx_paragraph(
                document, f"【你的错误作答】{item['wrong_answer']}",
                self._SIZE_ANALYSIS, align.LEFT, indent_cm=0.5,
            )

    # ------------------------------------------------------------------
    # PDF 渲染（PyMuPDF 直接排版，不依赖 Word 或外部打印工具）
    # ------------------------------------------------------------------

    _PDF_W = 595.0   # A4 宽（点）
    _PDF_H = 842.0   # A4 高（点）
    #: PDF 用：Unicode 上标 → 基字符（小号上移绘制）
    _SUPER_BASE = {
        "⁰": "0", "¹": "1", "²": "2", "³": "3", "⁴": "4", "⁵": "5",
        "⁶": "6", "⁷": "7", "⁸": "8", "⁹": "9",
        "⁺": "+", "⁻": "-", "⁼": "=", "⁽": "(", "⁾": ")",
        "ⁿ": "n", "ⁱ": "i",
        "ᵃ": "a", "ᵇ": "b", "ᶜ": "c", "ᵈ": "d", "ᵉ": "e", "ᶠ": "f",
        "ᵍ": "g", "ʰ": "h", "ʲ": "j", "ᵏ": "k", "ˡ": "l", "ᵐ": "m",
        "ᵒ": "o", "ᵖ": "p", "ʳ": "r", "ˢ": "s", "ᵗ": "t", "ᵘ": "u",
        "ᵛ": "v", "ʷ": "w", "ˣ": "x", "ʸ": "y", "ᶻ": "z",
    }
    #: PDF 用：Unicode 下标 → 基字符（小号下移绘制）
    _SUB_BASE = {
        "₀": "0", "₁": "1", "₂": "2", "₃": "3", "₄": "4",
        "₅": "5", "₆": "6", "₇": "7", "₈": "8", "₉": "9",
        "₊": "+", "₋": "-", "₌": "=", "₍": "(", "₎": ")",
        "ₐ": "a", "ₑ": "e", "ₕ": "h", "ᵢ": "i", "ⱼ": "j", "ₖ": "k",
        "ₗ": "l", "ₘ": "m", "ₙ": "n", "ₒ": "o", "ₚ": "p", "ᵣ": "r",
        "ₛ": "s", "ₜ": "t", "ᵤ": "u", "ᵥ": "v", "ₓ": "x",
    }
    #: 字体缺失字符的文本回退
    _TEXT_FALLBACK = {
        "·": "×", "\u2207": "grad", "ℝ": "R", "ℤ": "Z", "ℕ": "N", "ℂ": "C",
    }

    def _render_pdf(self, paper: Dict[str, Any], path: Path) -> None:
        try:
            import pymupdf
        except ImportError:  # pragma: no cover
            try:
                import fitz as pymupdf
            except ImportError as e:
                raise RuntimeError("缺少 PyMuPDF 依赖，无法导出 PDF 试卷") from e

        margin_h = self._MARGIN_H_CM * 28.35
        margin_v = self._MARGIN_V_CM * 28.35
        font_name, font_file = _find_pdf_font()
        font_obj = (
            pymupdf.Font(fontfile=font_file) if font_file else pymupdf.Font("china-s")
        )

        doc = pymupdf.open()
        state = {"y": margin_v + 12.0}

        def new_page() -> None:
            doc.new_page(width=self._PDF_W, height=self._PDF_H)
            if font_file:
                doc[-1].insert_font(fontname=font_name, fontfile=font_file)
            state["y"] = margin_v + 12.0

        def need_space(height: float) -> None:
            if state["y"] + height > self._PDF_H - margin_v - 16:
                new_page()

        def safe(text: str) -> str:
            for src, dst in self._TEXT_FALLBACK.items():
                if src in text:
                    text = text.replace(src, dst)
            return text

        def draw_mixed(x: float, baseline: float, text: str, size: float) -> None:
            """绘制文本，Unicode 上下标以小号偏移字符呈现。"""
            cursor = x
            index, total = 0, len(text)
            while index < total:
                ch = text[index]
                if ch in self._SUPER_BASE:
                    table, shift = self._SUPER_BASE, size * 0.32
                elif ch in self._SUB_BASE:
                    table, shift = self._SUB_BASE, -size * 0.22
                else:
                    table, shift = None, 0.0
                if table is not None:
                    end = index
                    while end < total and text[end] in table:
                        end += 1
                    segment = "".join(table[c] for c in text[index:end])
                    small = size * 0.65
                    doc[-1].insert_text(
                        (cursor, baseline - shift), segment,
                        fontname=font_name, fontfile=font_file, fontsize=small,
                    )
                    cursor += font_obj.text_length(segment, fontsize=small) + size * 0.05
                    index = end
                else:
                    end = index
                    while (
                        end < total
                        and text[end] not in self._SUPER_BASE
                        and text[end] not in self._SUB_BASE
                    ):
                        end += 1
                    segment = text[index:end]
                    doc[-1].insert_text(
                        (cursor, baseline), segment,
                        fontname=font_name, fontfile=font_file, fontsize=size,
                    )
                    cursor += font_obj.text_length(segment, fontsize=size)
                    index = end

        def wrap(text: str, size: float, max_width: float) -> List[str]:
            """按字体实际度量折行，保证中文行宽一致。"""
            lines: List[str] = []
            current = ""
            for ch in text:
                probe = current + ch
                if current and font_obj.text_length(probe, fontsize=size) > max_width:
                    lines.append(current)
                    current = ch
                else:
                    current = probe
            lines.append(current)
            return lines

        def write_line(text: str, size: float = 12.0, indent: float = 0.0) -> None:
            """写一段（按字体度量自动折行）。

            先剥掉公式标记，再把 ``\\frac`` 画成上下结构的真分数，
            其余内容转可读 Unicode，保证 PDF 里不出现原始 LaTeX 代码。
            """
            max_width = self._PDF_W - margin_h * 2 - indent
            for chunk, frac in split_fractions(strip_math_delimiters(text or "")):
                if chunk:
                    cleaned = safe(latex_to_unicode(chunk))
                    if not cleaned:
                        continue
                    for piece in wrap(cleaned, size, max_width):
                        need_space(size * 1.8)
                        draw_mixed(margin_h + indent, state["y"] + size, piece, size)
                        state["y"] += size * self._LINE_SPACING
                if frac:
                    self._pdf_fraction(
                        doc, state, need_space, draw_mixed, safe, frac, size,
                        indent, margin_h, font_name, font_file,
                    )

        def write_blank(lines: int, size: float = 12.0) -> None:
            """预留作答空白（每行含 1.5 倍行距）。"""
            for _ in range(max(0, lines)):
                need_space(size * self._LINE_SPACING)
                state["y"] += size * self._LINE_SPACING

        def write_centered(text: str, size: float) -> None:
            width = font_obj.text_length(text, fontsize=size)
            need_space(size * 2)
            draw_mixed(
                max(margin_h, (self._PDF_W - width) / 2),
                state["y"] + size,
                text,
                size,
            )
            state["y"] += size * 2

        # 标题（二号）+ 生成日期 / 知识点范围
        new_page()
        write_centered(paper["title"], self._SIZE_TITLE)
        write_centered(safe(paper["meta"]), self._SIZE_ANALYSIS)

        for index, section in enumerate(paper["sections"]):
            write_line(
                f"{_CN_NUMBERS[index]}、{section['label']}", size=self._SIZE_HEADING
            )
            state["y"] += 4
            for q in section["questions"]:
                write_line(f"{q['number']}．{q['stem']}", size=self._SIZE_QUESTION)
                for option in q.get("options") or []:
                    write_line(option, size=self._SIZE_QUESTION, indent=28)
                image = q.get("image")
                if image is not None:
                    self._pdf_image(doc, state, need_space, margin_h, image)
                if q.get("blank_lines"):
                    write_blank(int(q["blank_lines"]))
                state["y"] += 4

        if paper["answers"]:
            new_page()
            write_centered("参考答案与解析", self._SIZE_HEADING)
            state["y"] += 6
            for item in paper["answers"]:
                write_line(
                    f"{item['number']}．【答案】{item['answer']}",
                    size=self._SIZE_ANALYSIS,
                )
                if item.get("analysis"):
                    write_line(
                        f"【解析】{item['analysis']}",
                        size=self._SIZE_ANALYSIS, indent=16,
                    )
                if item.get("wrong_answer"):
                    write_line(
                        f"【你的错误作答】{item['wrong_answer']}",
                        size=self._SIZE_ANALYSIS, indent=16,
                    )
                state["y"] += 4

        self._pdf_decorations(doc, safe, font_name, font_file, font_obj, pymupdf)
        self._save_pdf(doc, path)

    @staticmethod
    def _save_pdf(doc, path: Path) -> None:
        """保存 PDF，并先把嵌入字体子集化。

        PyMuPDF 的 insert_font 会把**整个字体文件**塞进 PDF（SimSun 约 17MB），
        于是只有一道题的练习卷也能导出 18MB+，学生根本没法用微信/邮件发出去，
        打开也慢。子集化只保留实际用到的字形，实测单页卷 18.68MB → 0.05MB，
        页面显示与可复制文本完全一致。
        """
        try:
            doc.subset_fonts()
        except Exception as e:  # pragma: no cover - 旧版 PyMuPDF 无此方法
            logger.warning("PDF 字体子集化失败，导出文件会明显偏大: %s", e)
        doc.save(str(path), garbage=4, deflate=True)
        doc.close()

    def _pdf_fraction(
        self, doc, state, need_space, draw_mixed, safe, frac, size,
        indent, margin_h, font_name, font_file,
    ) -> None:
        """绘制上下结构分数：分子 / 横线 / 分母，与 Word 版式保持一致。"""
        num = safe(latex_to_unicode(frac[0]).strip()) or " "
        den = safe(latex_to_unicode(frac[1]).strip()) or " "
        small = max(7.0, size * 0.8)
        bar = "─" * max(2, (max(len(num), len(den)) + 1) // 2)
        need_space(small * 2 + 12)
        left = margin_h + indent + 6
        draw_mixed(left, state["y"] + small, num, small)
        doc[-1].insert_text(
            (left, state["y"] + small + 4), bar,
            fontname=font_name, fontfile=font_file, fontsize=small,
        )
        draw_mixed(left, state["y"] + small * 2 + 8, den, small)
        state["y"] += small * 2 + 14

    def _pdf_image(self, doc, state, need_space, margin_h, image: Path) -> None:
        """插入错题图片，按内容区宽度等比缩放并居中。"""
        import pymupdf

        try:
            with pymupdf.open(str(image)) as src:
                page_rect = src[0].rect
                width, height = page_rect.width, page_rect.height
        except Exception as e:
            logger.warning("Skip error image %s: %s", image, e)
            return
        if width <= 0 or height <= 0:
            return
        max_width = self._PDF_W - margin_h * 2
        scale = min(max_width / width, (self._PDF_H * 0.42) / height, 1.0)
        draw_width, draw_height = width * scale, height * scale
        need_space(draw_height + 10)
        left = (self._PDF_W - draw_width) / 2
        rect = pymupdf.Rect(left, state["y"], left + draw_width, state["y"] + draw_height)
        try:
            doc[-1].insert_image(rect, filename=str(image))
        except Exception as e:
            logger.warning("Insert error image %s failed: %s", image, e)
            return
        state["y"] += draw_height + 8

    def _pdf_decorations(
        self, doc, safe, font_name, font_file, font_obj, pymupdf
    ) -> None:
        """统一补页眉（含细分隔线）与页脚页码（此时才知道总页数）。"""
        total = len(doc)
        header = safe(PAPER_HEADER)
        size = 9.0
        for index, page in enumerate(doc, start=1):
            width = font_obj.text_length(header, fontsize=size)
            page.insert_text(
                ((self._PDF_W - width) / 2, 46), header,
                fontname=font_name, fontfile=font_file, fontsize=size,
            )
            page.draw_line(
                pymupdf.Point(60, 56), pymupdf.Point(self._PDF_W - 60, 56),
                color=(0.65, 0.65, 0.65), width=0.6,
            )
            footer = f"第 {index} 页 / 共 {total} 页"
            width = font_obj.text_length(footer, fontsize=size)
            page.insert_text(
                ((self._PDF_W - width) / 2, self._PDF_H - 34), footer,
                fontname=font_name, fontfile=font_file, fontsize=size,
            )


def _find_pdf_font() -> Tuple[str, Optional[str]]:
    """返回 (fontname, fontfile)：优先系统中文宋体，找不到回退内置 china-s。"""
    for candidate in (
        "C:/Windows/Fonts/simsun.ttc",
        "C:/Windows/Fonts/msyh.ttc",
        "C:/Windows/Fonts/msyh.ttf",
        "C:/Windows/Fonts/simhei.ttf",
    ):
        if Path(candidate).exists():
            return "paperfont", candidate
    return "china-s", None
